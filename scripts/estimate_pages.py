# -*- coding: utf-8 -*-
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

f = r"c:\Users\ASUS\projects\ai-document-bot\Report_short_35p.docx"
d = Document(f)


def iter_block_items(parent):
    for child in parent.element.body.iterchildren():
        if child.tag.endswith('}p'):
            yield Paragraph(child, parent)
        elif child.tag.endswith('}tbl'):
            yield Table(child, parent)


def table_words(t):
    w = 0
    for row in t.rows:
        for c in row.cells:
            w += len(c.text.split())
    return w


main_words = 0
in_main = False
n_main_tables = 0
n_fig_placeholders = 0
headings = 0

for blk in iter_block_items(d):
    if isinstance(blk, Paragraph):
        txt = blk.text.strip()
        style = blk.style.name if blk.style else ''
        if txt.startswith('Chapter 1'):
            in_main = True
        if txt == 'Bibliography':
            in_main = False
        if not in_main:
            continue
        if style in ('Heading 1', 'Heading 2', 'Heading 3'):
            headings += 1
        if txt.startswith('[ PASTE'):
            n_fig_placeholders += 1
            continue
        main_words += len(txt.split())
    else:
        if in_main:
            n_main_tables += 1
            main_words += table_words(blk)

print('Main-content words (incl. table text):', main_words)
print('Main-content tables:', n_main_tables)
print('Figure placeholders (become images):', n_fig_placeholders)
print('Headings:', headings)

prose_pages = main_words / 340.0
table_extra = n_main_tables * 0.25
fig_extra = n_fig_placeholders * 0.5
heading_extra = headings * 0.04
est = prose_pages + table_extra + fig_extra + heading_extra

print()
print('Rough prose pages:        %.1f' % prose_pages)
print('Table vertical allowance: %.1f' % table_extra)
print('Figure image allowance:   %.1f' % fig_extra)
print('Heading/spacing:          %.1f' % heading_extra)
print('ESTIMATED MAIN-CONTENT PAGES: %.0f' % round(est))
